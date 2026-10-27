import json,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
d=json.load(open('content.json'))
def rng(s):
    s=str(s)
    if s.count('–')>=2:
        p=s.split('–'); return f"{p[0]} → {p[-1]}"
    return s
def stats(js):
    o=json.loads(js) if js else {}
    return '；'.join(f"{k} {rng(v)}" for k,v in o.items())
CAT={'melee':'近战','bow':'弩','caster':'法器','gun':'铳','shield':'盾'}
wb=Workbook(); ws=wb.active; ws.title='说明'
for r in [["灰烬地窖 · 现有内容总表"],["导出自 main 105540b 的游戏运行时数据（2026-10-26）。区间“a → b”= 品质 0 → 品质 10。"],
 ["页签：武器底子 / 其他底子 / 技能石 / 辅助符文 / 词条 / 独特 / 誓约 / 铸符"],["权重 = 掉落权重；0 表示不再掉落。词条页只列现在还能掷出的词条。"]]: ws.append(r)
ws['A1'].font=Font(bold=True,size=14); ws.column_dimensions['A'].width=110
def sheet(name,head,rows,widths):
    s=wb.create_sheet(name); s.append(head)
    for c in s[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='4A3F55')
    for r in rows: s.append(r)
    for i,w in enumerate(widths): s.column_dimensions[get_column_letter(i+1)].width=w
    for row in s.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    s.freeze_panes='C2'; s.auto_filter.ref=s.dimensions
sheet('武器底子',['id','名称','类别','元素','对应技能石','需求等级','权重','最小伤害','最大伤害','攻速','暴击%','射程','已退役','标签'],
 [[w['id'],w['n'],CAT.get(w['cat'],w['cat']),w['el'],w['gem'],w['req'],w['w'],rng(w['min']),rng(w['max']),rng(w['aps']),rng(w['crit']),rng(w['range']),w['retired'],w['tags']] for w in d['weapons']],
 [16,12,8,8,12,8,8,14,14,8,8,8,8,50])
sheet('其他底子',['部位','id','名称','需求等级','权重','定位','隐性支援（副手）','基础数值（品质 0 → 10）'],
 [[a['kind'],a['id'],a['n'],a['req'],a['w'],a['d'],a['sup'],stats(a['stats'])] for a in d['armors']],[8,18,12,8,8,26,40,70])
sheet('技能石',['id','名称','形态','元素','标签','倍率 mul','速度 spd','需求等级','权重','AI','来自旧底子','并入的形态符文','特殊效果字段'],
 [[g['id'],g['n'],g['form'],g['el'],g['tags'],g['mul'],g['spd'],g['req'],g['w'],g['ai'],g['src'],g['fr'],g['fx']] for g in d['gems']],[14,10,10,10,40,10,8,8,8,8,24,14,30])
def fit(f): return re.sub(r'tag:(\w+)',lambda m:m.group(1),f)
order={'辅助':0,'形态（已并进技能石）':1}
sheet('辅助符文',['id','名称','类型','可插在','需求等级','权重','效果（1 级）','效果（满级）'],
 [[m['id'],m['n'],m['type'],fit(m['fit']),m['req'],m['w'],m['d1'],m['dm']] for m in sorted(d['mods'],key=lambda m:(order.get(m['type'],2),m['req']))],[14,12,14,16,8,8,60,60])
sheet('词条',['key','前/后缀','可出部位','标签限制','元素','档位数','最低档','最高档','最高档物品等级','总权重'],
 [[a['k'],a['type'],a['on'],a['tag'],a['el'],a['tiers'],a['t1'],a['top'],a['topIlvl'],a['w']] for a in sorted(d['affs'],key=lambda a:(a['on'],a['type']))],[24,8,26,12,8,8,30,30,10,8])
sheet('独特',['id','名称','部位','底子','需求等级','权重','掉落限制','固定词条','规则'],
 [[u['id'],u['n'],u['kind'],u['base'],u['req'],u['w'],u['ladder'],u['fixed'],u['rules']] for u in sorted(d['uniq'],key=lambda u:(u['kind'],u['req']))],[16,12,8,12,8,8,10,26,90])
sheet('誓约',['id','可出部位','正面','代价'],[[o['id'],o['on'],o['pos'],o['neg']] for o in d['oath']],[22,20,40,40])
sheet('铸符',['id','名称','说明'],[[g['id'],g['n'],g['d']] for g in d['glyph']],[12,12,60])
wb.save('/mnt/project-files/design/tables/现有内容总表.xlsx')
print({k:len(v) for k,v in d.items()})
